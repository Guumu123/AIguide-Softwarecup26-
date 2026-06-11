import asyncio
import os
import sys
import pandas as pd
from docx import Document
from sqlalchemy.ext.asyncio import AsyncSession

# Set up path to allow importing from 'app'
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.db.session import AsyncSessionLocal
from app.models.base_models import Attraction
from app.models.admin_models import TouristBehavior, KBDocument

DATA_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "知识库", "示范景区公开资料包")

async def import_attractions(session: AsyncSession):
    """从 docx 表格中导入景点结构化数据"""
    path = os.path.join(DATA_DIR, "灵山胜境 景点结构化数据集.docx")
    if not os.path.exists(path):
        print(f"File not found: {path}")
        return

    doc = Document(path)
    if not doc.tables:
        print("No tables found in attractions docx.")
        return

    table = doc.tables[0]
    # 获取表头
    headers = [cell.text.strip() for cell in table.rows[0].cells]
    
    # 建立中英文映射
    mapping = {
        "景点ID": "attraction_id",
        "景点名称": "name",
        "所属景区": "scenic_area",
        "地理位置": "location",
        "主要参数": "parameters",
        "核心功能": "core_function",
        "文化内涵": "cultural_meaning",
        "详细介绍": "detailed_intro",
        "亮点介绍": "highlights",
        "演出信息": "performance_info",
        "备注": "notes",
        "预计游览时长": "estimated_duration",
        "类别": "category"
    }

    count = 0
    for row in table.rows[1:]:
        row_data = {}
        for i, cell in enumerate(row.cells):
            ch_header = headers[i]
            en_field = mapping.get(ch_header)
            if en_field:
                val = cell.text.strip()
                if en_field == "estimated_duration":
                    try:
                        val = float(val) if val else 0.0
                    except ValueError:
                        val = 0.0
                row_data[en_field] = val
        
        if row_data.get("attraction_id"):
            attraction = Attraction(**row_data)
            session.add(attraction)
            count += 1
    
    await session.commit()
    print(f"Successfully imported {count} attractions.")

async def import_tourist_behaviors(session: AsyncSession):
    """从 xlsx 中导入游客行为分析数据"""
    path = os.path.join(DATA_DIR, "景点景区旅游数据行为分析数据.xlsx")
    if not os.path.exists(path):
        print(f"File not found: {path}")
        return

    df = pd.read_excel(path)
    
    mapping = {
        '游客ID': 'tourist_id',
        '用户昵称': 'user_nickname',
        '年龄': 'age',
        '性别': 'gender',
        '景点名称': 'attraction_name',
        '景点类型': 'attraction_type',
        '游览日期': 'visit_date',
        '停留时长': 'stay_duration',
        '门票花费': 'ticket_cost',
        '餐饮花费': 'food_cost',
        '购物花费': 'shopping_cost',
        '交通花费': 'transport_cost',
        '娱乐花费': 'entertainment_cost',
        '总花费': 'total_cost',
        '团队人数': 'group_size',
        '满意度': 'satisfaction'
    }
    
    batch_size = 1000
    total_count = 0
    for i in range(0, len(df), batch_size):
        batch = df.iloc[i:i+batch_size]
        records = []
        for _, row in batch.iterrows():
            record_data = {}
            for ch_col, en_col in mapping.items():
                if ch_col in row:
                    val = row[ch_col]
                    # Handle NaN values from pandas
                    if pd.isna(val):
                        val = None
                    record_data[en_col] = val
            records.append(TouristBehavior(**record_data))
        
        session.add_all(records)
        await session.commit()
        total_count += len(records)
    
    print(f"Successfully imported {total_count} tourist behavior records.")

async def import_kb_documents(session: AsyncSession):
    """导入非结构化知识库文档"""
    path = os.path.join(DATA_DIR, "灵山胜境：历史、文化、景点特色与个性化游览指南.docx")
    if not os.path.exists(path):
        print(f"File not found: {path}")
        return

    doc = Document(path)
    content = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    
    kb_doc = KBDocument(
        title="灵山胜境：历史、文化、景点特色与个性化游览指南",
        content=content,
        doc_type="讲解词",
        source_file="灵山胜境：历史、文化、景点特色与个性化游览指南.docx",
        is_active=True
    )
    session.add(kb_doc)
    await session.commit()
    print("Successfully imported KB document.")

async def main():
    async with AsyncSessionLocal() as session:
        print("Starting data import...")
        try:
            await import_attractions(session)
            await import_tourist_behaviors(session)
            await import_kb_documents(session)
            print("Data import completed successfully.")
        except Exception as e:
            print(f"Error during import: {e}")
            await session.rollback()

if __name__ == "__main__":
    asyncio.run(main())
